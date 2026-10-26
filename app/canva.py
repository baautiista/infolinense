"""Canva OAuth and brand-template autofill. Tokens remain on the server volume."""
import base64
import hashlib
import json
import re
import os
import secrets
import threading
import time
from pathlib import Path
from urllib.parse import urlencode

import requests

from . import db, layout
from .config import PUBLIC_BASE_URL, RENDER_DIR

CLIENT_ID = os.getenv('CANVA_CLIENT_ID', '').strip()
CLIENT_SECRET = os.getenv('CANVA_CLIENT_SECRET', '').strip()
TEMPLATE_ID = os.getenv('CANVA_BRAND_TEMPLATE_ID', '').strip()
# Variantes de la plantilla según las líneas del titular: la entradilla queda fija abajo,
# el titular crece hacia arriba y la sección se coloca justo encima.
_DEFAULT_VARIANTS = {'EAHWTjWEEnA': {1: 'EAHXJaqiN5w', 2: 'EAHXJVlUZ2g', 4: 'EAHXJyNUgnU'}}
TEMPLATE_BY_LINES = {
    1: os.getenv('CANVA_TEMPLATE_1_LINE', '').strip() or _DEFAULT_VARIANTS.get(TEMPLATE_ID, {}).get(1, ''),
    2: os.getenv('CANVA_TEMPLATE_2_LINES', '').strip() or _DEFAULT_VARIANTS.get(TEMPLATE_ID, {}).get(2, ''),
    4: os.getenv('CANVA_TEMPLATE_4_LINES', '').strip() or _DEFAULT_VARIANTS.get(TEMPLATE_ID, {}).get(4, ''),
}


def _brand_templates(brand):
    """Plantillas del medio: {main, 1, 2, 4}. InfoLinense usa las de Railway si no se han cambiado en Ajustes."""
    from . import brands
    t = brands.settings(brand)['templates']
    if brands.valid(brand) == brands.DEFAULT:
        return {'main': t.get('main') or TEMPLATE_ID, '1': t.get('1') or TEMPLATE_BY_LINES.get(1, ''),
                '2': t.get('2') or TEMPLATE_BY_LINES.get(2, ''), '4': t.get('4') or TEMPLATE_BY_LINES.get(4, '')}
    return t


def template_for(headline, brand='infolinense'):
    t = _brand_templates(brand)
    lines = len(layout.headline_lines(headline))
    return t.get(str(lines)) or t.get('main') or ''
API = 'https://api.canva.com/rest/v1'
SCOPES = 'asset:read asset:write brandtemplate:content:read design:content:read design:content:write design:meta:read'
_lock = threading.RLock()


def ready():
    return bool(CLIENT_ID and CLIENT_SECRET and TEMPLATE_ID and PUBLIC_BASE_URL)


def connected():
    return db.row('SELECT 1 FROM canva_oauth WHERE id=1') is not None


def callback_url():
    return PUBLIC_BASE_URL + '/api/canva/callback'


def authorization_url():
    if not ready():
        raise ValueError('Configura CANVA_CLIENT_ID, CANVA_CLIENT_SECRET y CANVA_BRAND_TEMPLATE_ID')
    state = secrets.token_urlsafe(40)
    verifier = secrets.token_urlsafe(64)
    challenge = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b'=').decode()
    db.exec_('DELETE FROM canva_oauth_state WHERE expires_at<?', (int(time.time()),))
    db.exec_('INSERT INTO canva_oauth_state VALUES(?,?,?)', (state, verifier, int(time.time()) + 600))
    return 'https://www.canva.com/api/oauth/authorize?' + urlencode({
        'code_challenge': challenge, 'code_challenge_method': 'S256',
        'scope': SCOPES, 'response_type': 'code', 'client_id': CLIENT_ID,
        'state': state, 'redirect_uri': callback_url(),
    })


def _token_request(form):
    resp = requests.post(API + '/oauth/token', data=form,
                         auth=(CLIENT_ID, CLIENT_SECRET), timeout=25)
    if not resp.ok:
        raise ValueError('Canva rechazó la autorización (%s)' % resp.status_code)
    data = resp.json()
    if not data.get('access_token') or not data.get('refresh_token'):
        raise ValueError('Canva no devolvió los tokens necesarios')
    return data


def _save_token(data):
    with _lock:
        db.exec_('INSERT OR REPLACE INTO canva_oauth VALUES(1,?,?,?)',
                 (data['access_token'], data['refresh_token'],
                  int(time.time()) + int(data.get('expires_in', 3600))))


def complete(code, state):
    row = db.row('SELECT verifier,expires_at FROM canva_oauth_state WHERE state=?', (state,))
    if not row or row['expires_at'] < int(time.time()):
        raise ValueError('Autorización caducada o no iniciada desde InfoLinense Desk')
    db.exec_('DELETE FROM canva_oauth_state WHERE state=?', (state,))
    data = _token_request({'grant_type': 'authorization_code', 'code': code,
                           'code_verifier': row['verifier'], 'redirect_uri': callback_url()})
    _save_token(data)


def access_token():
    with _lock:
        row = db.row('SELECT * FROM canva_oauth WHERE id=1')
        if not row:
            raise ValueError('Conecta tu cuenta de Canva desde Ajustes')
        if row['expires_at'] > time.time() + 120:
            return row['access_token']
        data = _token_request({'grant_type': 'refresh_token', 'refresh_token': row['refresh_token']})
        _save_token(data)
        return data['access_token']


def _api(method, path, **kwargs):
    try:
        resp = requests.request(method, API + path,
                                headers={'Authorization': 'Bearer ' + access_token(), **kwargs.pop('headers', {})},
                                timeout=45, **kwargs)
        if not resp.ok:
            payload = resp.json()
            detail = payload.get('error') or payload
            if isinstance(detail, dict):
                detail = detail.get('message') or detail.get('code')
            raise ValueError('Canva (%s): %s' % (resp.status_code, str(detail or 'revisa permisos y configuración')[:180]))
        return resp.json()
    except requests.RequestException as exc:
        raise ValueError('No se pudo contactar con Canva') from exc


def _wait(path, first=None):
    job = (first or {}).get('job') or {}
    for attempt in range(40):
        if attempt or not job.get('status'):
            job = _api('GET', path).get('job', {})
        if job.get('status') == 'success':
            return job
        if job.get('status') == 'failed':
            err = job.get('error') or {}
            detail = err.get('message') or err.get('code') if isinstance(err, dict) else ''
            raise ValueError('Canva no pudo completar el trabajo' + (': ' + str(detail)[:180] if detail else ''))
        time.sleep(2)
    raise ValueError('Canva aún procesa el diseño; vuelve a intentarlo en unos segundos')


def _design_from_job(job):
    """Canva devuelve el diseño en job.result.design (versiones antiguas: job.design)."""
    result = job.get('result') or {}
    design = result.get('design') or job.get('design') or {}
    return design if isinstance(design, dict) else {}


def template_status():
    dataset = _api('GET', '/brand-templates/' + TEMPLATE_ID + '/dataset').get('dataset', {})
    expected = {'HEADLINE': 'text', 'SUMMARY': 'text', 'SECTION': 'text', 'PHOTO': 'image'}
    return {'valid': all(dataset.get(k, {}).get('type') == v for k, v in expected.items()),
            'fields': {k: dataset.get(k, {}).get('type') for k in expected}}


def permission_status():
    if not connected():
        return {'connected': False, 'can_export_png': False}
    try:
        response = requests.post(API + '/oauth/introspect',
                                 auth=(CLIENT_ID, CLIENT_SECRET),
                                 data={'token': access_token()}, timeout=20)
        response.raise_for_status()
        payload = response.json()
        scopes = set((payload.get('scope') or '').split())
        return {'connected': bool(payload.get('active')),
                'can_export_png': 'design:content:read' in scopes}
    except requests.RequestException as exc:
        raise ValueError('No se pudieron comprobar los permisos de Canva') from exc


def _content_hash(article):
    fields = [article.get(k) or '' for k in ('headline', 'subtitle', 'section', 'image_url')]
    headline_size = article.get('headline_size') or 'auto'
    photo_zoom = float(article.get('photo_zoom') or 1)
    photo_x = float(article.get('photo_x') if article.get('photo_x') is not None else 50)
    photo_y = float(article.get('photo_y') if article.get('photo_y') is not None else 50)
    if headline_size != 'auto' or abs(photo_zoom - 1) > 0.001 or abs(photo_x - 50) > 0.001 or abs(photo_y - 50) > 0.001:
        fields += [headline_size, photo_zoom, photo_x, photo_y]
    if article.get('section_label') or article.get('brand') not in (None, '', 'infolinense'):
        fields += [article.get('section_label') or '', article.get('brand') or '']
    return hashlib.sha256(json.dumps(fields, ensure_ascii=False).encode()).hexdigest()


def _is_default(article):
    from . import brands
    return brands.valid(article.get('brand')) == brands.DEFAULT


def _page_for(article):
    if not _is_default(article):  # otros medios: página de su sección o la indicada en Ajustes → Medios
        from . import brands
        return brands.section_page(article.get('brand'), article.get('section'), bool(article.get('_carousel')), article.get('headline') or '')
    return layout.family_for(article.get('section'), article.get('headline'))[1]


def design_fields(article, strict=True):
    """Textos de la plantilla: exactamente el titular y la entradilla escritos, completos."""
    family = layout.family_for(article.get('section'), article.get('headline'))
    headline = re.sub(r'\s+', ' ', article.get('headline') or '').strip()
    summary = re.sub(r'\s+', ' ', article.get('subtitle') or '').strip()
    # Titular y entradilla van ENTEROS, tal cual se han escrito: nunca se recortan ni se acortan.
    if not _is_default(article):
        from . import brands
        b = brands.settings(article.get('brand'))
        sec = brands.find_section(b['slug'], article.get('section')) \
            or brands.detect_section(b['slug'], article.get('headline') or '') or (article.get('section') or '').strip()
        label = sec
        if b.get('free_title_section') and sec == b['free_title_section']:
            label = (article.get('section_label') or '').strip() or brands.holy_week_label()  # «General» con título propio
        return {'HEADLINE': headline, 'SUMMARY': summary, 'SECTION': label,
                'page': _page_for(article), 'template': template_for(headline, article.get('brand'))}
    return {'HEADLINE': headline, 'SUMMARY': summary, 'SECTION': layout.normalize_section(article.get('section'), article.get('headline') or ''),
            'page': family[1], 'template': template_for(headline)}


# Tamaño de letra según la longitud: la plantilla puede tener, en el mismo sitio, cuadros con letra
# más grande (…_L) y más pequeña (…_S). Se rellena el que corresponde y los demás se dejan en blanco.
SIZE_RULES = {'HEADLINE': (40, 80), 'SECTION': (12, 18)}   # hasta N1 caracteres: grande; más de N2: pequeña
BLANK = ' '


def size_variants(schema, texts, headline_size='auto'):
    out = {}
    for key, (big, small) in SIZE_RULES.items():
        text = texts.get(key) or ''
        has_l = schema.get(key + '_L', {}).get('type') == 'text'
        has_s = schema.get(key + '_S', {}).get('type') == 'text'
        if not (has_l or has_s) or not text:
            continue
        has_n = schema.get(key, {}).get('type') == 'text'
        mode = str(headline_size or 'auto').lower() if key == 'HEADLINE' else 'auto'
        if mode == 'large' and has_l:
            pick = key + '_L'
        elif mode == 'small' and has_s:
            pick = key + '_S'
        else:
            pick = key + '_L' if has_l and len(text) <= big else key + '_S' if has_s and len(text) > small else key if has_n else key + ('_S' if has_s else '_L')
        for k in (key, key + '_L', key + '_S'):
            if schema.get(k, {}).get('type') == 'text':
                out[k] = {'type': 'text', 'text': text if k == pick else BLANK}
    return out


def _photo_for_canvas(photo, article):
    """Prepara la foto completa para que el encuadre se pueda ajustar dentro de Canva."""
    from io import BytesIO
    from PIL import Image, ImageOps

    with Image.open(photo) as source:
        image = ImageOps.exif_transpose(source).convert('RGB')
    # Limita la resolución sin recortar la foto; Canva recibe toda la imagen y
    # permite moverla/ajustar su tamaño en el marco de la plantilla original.
    resampling = getattr(Image, 'Resampling', Image).LANCZOS
    image.thumbnail((3000, 3000), resampling)
    buffer = BytesIO()
    image.save(buffer, format='JPEG', quality=90, optimize=True)
    return buffer.getvalue()


def _page_count(design_id):
    try:
        return int((_api('GET', '/designs/' + design_id).get('design') or {}).get('page_count') or 0)
    except Exception:
        return 0


def _keep_only_page(design_id, page):
    """Deja en el diseño solo la página de la noticia (la plantilla trae todas las secciones).
    Así, al abrirlo en Canva para retocarlo, solo hay una página. Si Canva no lo permite, se deja como estaba."""
    total = _page_count(design_id)
    page = int(page or 1)
    if total <= 1 or page > total:
        return total <= 1
    others = [n for n in range(1, total + 1) if n != page]
    try:
        job = _api('POST', '/merges', json={'type': 'modify_existing_design', 'design_id': design_id,
                                             'operations': [{'type': 'delete_pages', 'page_numbers': others}]})
        _wait('/merges/' + job['job']['id'], job)
        return True
    except Exception:
        return False


def _export_png(article, design_id, output_suffix=None, mark_main=True, page=None):
    from io import BytesIO
    from PIL import Image
    from urllib.parse import urlparse

    export = _api('POST', '/exports', json={'design_id': design_id,
                  'format': {'type': 'png', 'pages': [int(page or _page_for(article))]}})
    urls = _wait('/exports/' + export['job']['id']).get('urls') or []
    if not urls:
        raise ValueError('Canva no devolvió la imagen exportada')
    host = urlparse(urls[0]).hostname or ''
    if urlparse(urls[0]).scheme != 'https' or not (host == 'canva.com' or host.endswith('.canva.com')):
        raise ValueError('Canva devolvió una URL de descarga inesperada')
    try:
        response = requests.get(urls[0], timeout=45)
        response.raise_for_status()
        if len(response.content) > 20_000_000:
            raise ValueError('La imagen exportada supera 20 MB')
        image = Image.open(BytesIO(response.content))
        if image.format != 'PNG':
            raise ValueError('Canva no devolvió un archivo PNG')
        image.verify()
        safe_suffix = ''.join(ch for ch in str(output_suffix or '') if ch.isalnum() or ch in '_-')
        filename = 'article_%s%s.png' % (article['id'], ('_' + safe_suffix) if safe_suffix else '')
        output = RENDER_DIR / filename
        temp = output.with_suffix('.tmp')
        temp.write_bytes(response.content)
        temp.replace(output)
    except (requests.RequestException, OSError) as exc:
        raise ValueError('No se pudo descargar o validar el PNG de Canva') from exc
    if mark_main:
        db.exec_('UPDATE articles SET render_path=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',
                 (str(output), article['id']))
        db.exec_('UPDATE canva_designs SET exported=1,updated_at=CURRENT_TIMESTAMP WHERE article_id=?',
                 (article['id'],))
    return str(output)

def export_design(article):
    if not ready() or not connected():
        raise ValueError('Conecta Canva antes de exportar el diseño')
    design = db.row('SELECT * FROM canva_designs WHERE article_id=?', (article['id'],))
    if not design:
        raise ValueError('Primero crea el diseño con la plantilla de Canva')
    if design.get('content_hash') != _content_hash(article):
        raise ValueError('La noticia o la foto cambió; vuelve a crear el diseño de Canva')
    db.exec_('UPDATE canva_designs SET exported=0 WHERE article_id=?', (article['id'],))
    db.exec_('UPDATE articles SET render_path=NULL,status=CASE WHEN status="approved" THEN "review_ready" ELSE status END WHERE id=?', (article['id'],))
    _export_png(article, design['design_id'], page=1 if _page_count(design['design_id']) == 1 else _page_for(article))
    return {'url': design['url'], 'design_id': design['design_id'], 'exported': True}


def create_design(article, store=True, title_suffix='', output_suffix=None, strict=True):
    if not ready() or not connected():
        raise ValueError('Conecta tu cuenta de Canva desde Ajustes')
    photo = Path(article.get('image_local') or '')
    if not photo.is_file():
        raise ValueError('Elige primero una foto para la noticia')
    from . import brands
    brand = brands.settings(article.get('brand'))
    texts = design_fields(article, strict)
    template_id = texts['template']
    if not template_id:
        raise ValueError('Falta la plantilla de Canva de %s. Pégala en Ajustes → Medios.' % brand['name'])
    schema = _api('GET', '/brand-templates/' + template_id + '/dataset').get('dataset', {})
    expected = {'HEADLINE': 'text', 'PHOTO': 'image'}
    if _is_default(article):
        expected.update(SECTION='text', SUMMARY='text')
    missing = [key for key, kind in expected.items() if schema.get(key, {}).get('type') != kind]
    if missing:
        raise ValueError('A la plantilla de Canva de %s le faltan los campos de datos: %s. En Canva, en la plantilla, '
                         'pon esos nombres a los cuadros de texto y a la foto (Aplicaciones → Autocompletar).' % (brand['name'], ', '.join(missing)))
    data = _photo_for_canvas(photo, article)
    if len(data) > 20_000_000:
        raise ValueError('La fotografía supera 20 MB')
    suffix = ''.join(ch for ch in str(output_suffix or '') if ch.isalnum() or ch in '_-')
    filename = re.sub(r'\W+', '', brand['name']) + '-' + str(article['id']) + (('-' + suffix) if suffix else '') + '.jpg'
    metadata = json.dumps({'name_base64': base64.b64encode(filename.encode()).decode()})
    upload = _api('POST', '/asset-uploads', data=data,
                  headers={'Content-Type': 'application/octet-stream', 'Asset-Upload-Metadata': metadata})
    done = _wait('/asset-uploads/' + upload['job']['id'], upload)
    asset = done.get('asset') or (done.get('result') or {}).get('asset') or {}
    if not asset.get('id'):
        raise ValueError('Canva no devolvió la fotografía cargada')
    fields = {
        'HEADLINE': {'type': 'text', 'text': texts['HEADLINE']},
        'SUMMARY': {'type': 'text', 'text': texts['SUMMARY']},
        'SECTION': {'type': 'text', 'text': texts['SECTION']},
        'PHOTO': {'type': 'image', 'asset_id': asset['id']},
    }
    for opt in ('SECTION', 'SUMMARY'):  # plantillas sin sección o sin entradilla
        if schema.get(opt, {}).get('type') != 'text' or not texts[opt]:
            fields.pop(opt)
    fields.update(size_variants(schema, texts, 'auto'))
    name = brand['name'] + ' · ' + (article.get('headline') or '')[:100]
    if title_suffix:
        name += ' · ' + str(title_suffix)[:40]
    job = _api('POST', '/autofills', json={'type': 'create_from_brand_template',
                    'brand_template_id': template_id,
                    'title': name, 'data': fields})
    done = _wait('/autofills/' + job['job']['id'], job)
    design = _design_from_job(done)
    if not design.get('id'):
        raise ValueError('Canva terminó el trabajo pero no devolvió el diseño (respuesta: %s)' % json.dumps(done, ensure_ascii=False)[:200])
    # design.url es el enlace estable de edición; urls.edit_url caduca a los 30 días.
    url = design.get('url') or (design.get('urls') or {}).get('edit_url') or (design.get('urls') or {}).get('view_url') \
        or 'https://www.canva.com/design/%s/edit' % design['id']
    output = {'url': url, 'design_id': design.get('id'), 'exported': False}
    if store:
        db.exec_('INSERT OR REPLACE INTO canva_designs(article_id,design_id,url,exported,content_hash,updated_at) VALUES(?,?,?,0,?,CURRENT_TIMESTAMP)',
                 (article['id'], design['id'], url, _content_hash(article)))
        db.exec_('UPDATE articles SET render_path=NULL,status=CASE WHEN status="approved" THEN "review_ready" ELSE status END WHERE id=?', (article['id'],))
    page = 1 if _keep_only_page(design['id'], texts['page']) else texts['page']
    try:
        _export_png(article, design['id'], output_suffix=output_suffix, mark_main=store, page=page)
        output['exported'] = True
    except (ValueError, requests.RequestException, KeyError, OSError) as exc:
        output['export_error'] = str(exc)
    return output


def create_carousel_designs(article, slides, selected_photos):
    from . import photos as photo_service

    if not 3 <= len(slides) <= 6:
        raise ValueError('El carrusel debe tener entre tres y seis diapositivas')
    if not 1 <= len(selected_photos) <= 6:
        raise ValueError('Selecciona entre una y seis fotografías autorizadas')
    if len(selected_photos) < len(slides):
        selected_photos = [selected_photos[index % len(selected_photos)] for index in range(len(slides))]
    elif len(selected_photos) > len(slides):
        selected_photos = selected_photos[:len(slides)]
    try:
        allowed = json.loads(article.get('image_candidates_json') or '[]')
    except Exception:
        allowed = []

    for old in RENDER_DIR.glob('article_%s_carousel_*.png' % article['id']):
        try:
            old.unlink()
        except OSError:
            pass
    db.exec_('DELETE FROM carousel_designs WHERE article_id=?', (article['id'],))

    results = []
    for index, (slide, photo) in enumerate(zip(slides, selected_photos), start=1):
        try:
            local_path = photo_service.download_image(photo['url'])
        except Exception as exc:
            raise ValueError('No se pudo descargar la foto de la diapositiva %s' % index) from exc
        page = dict(article)
        page.update({
            '_carousel': True,
            'headline': slide.get('title') or '',
            'image_headline': '',
            'graphic_summary': slide.get('text') or '',
            'subtitle': slide.get('text') or '',
            'image_url': photo.get('url') or '',
            'image_source': photo.get('source') or '',
            'image_license': photo.get('license') or '',
            'image_local': local_path,
            'image_candidates_json': json.dumps([photo], ensure_ascii=False),
        })
        suffix = 'carousel_%s' % index
        result = create_design(
            page, store=False, title_suffix='Carrusel %s/%s' % (index, len(slides)),
            output_suffix=suffix, strict=False
        )
        if not result.get('exported'):
            raise ValueError(result.get('export_error') or 'Canva no exportó la diapositiva %s' % index)
        db.exec_(
            '''INSERT OR REPLACE INTO carousel_designs
               (article_id,slide_index,design_id,url,exported,content_hash,image_url,image_source,image_license,updated_at)
               VALUES(?,?,?,?,1,?,?,?,?,CURRENT_TIMESTAMP)''',
            (article['id'], index, result['design_id'], result['url'], _content_hash(page),
             photo.get('url') or '', photo.get('source') or '', photo.get('license') or '')
        )
        results.append({
            'index': index, 'title': slide.get('title') or '',
            'text': slide.get('text') or '', 'url': result['url'],
            'exported': True, 'image_source': photo.get('source') or '',
            'image_license': photo.get('license') or '',
            'download_path': '/api/articles/%s/carousel/slide/%s' % (article['id'], index),
        })
    first_url = results[0]['url'] if results else None
    return {
        'slides': results,
        'url': first_url,
        'design_url': first_url,
        'download_url': '/api/articles/%s/carousel/download' % article['id'],
    }


def _autofill_text_only(article):
    """Diseño de la noticia con titular, entradilla y sección (sin foto). Devuelve (design_id, página)."""
    texts = design_fields(article, strict=False)
    template_id = texts['template']
    if not template_id:
        raise ValueError('Falta la plantilla de Canva del medio')
    schema = _api('GET', '/brand-templates/' + template_id + '/dataset').get('dataset', {})
    fields = {k: {'type': 'text', 'text': texts[k]} for k in ('HEADLINE', 'SUMMARY', 'SECTION')
              if schema.get(k, {}).get('type') == 'text' and texts.get(k)}
    fields.update(size_variants(schema, texts, 'auto'))
    job = _api('POST', '/autofills', json={'type': 'create_from_brand_template', 'brand_template_id': template_id,
                                          'title': 'Lote · ' + (article.get('headline') or '')[:80], 'data': fields})
    design = _design_from_job(_wait('/autofills/' + job['job']['id'], job))
    if not design.get('id'):
        raise ValueError('Canva no devolvió el diseño')
    return design['id'], int(texts['page'] or 1)


def _batch_pptx(articles):
    """Presentación 1080×1350 con una página por noticia: etiqueta de sección, titular y entradilla ENTEROS
    y un hueco para la foto. Se importa en Canva como UN solo diseño editable."""
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from . import brands
    px = lambda v: Emu(int(v * 9525))
    rgb = lambda h: RGBColor.from_string(h.lstrip('#').upper())
    prs = Presentation(); prs.slide_width, prs.slide_height = px(1080), px(1350)
    blank = prs.slide_layouts[6]
    for a in articles:
        brand = brands.settings(a.get('brand'))
        fam = layout.family_for(a.get('section'), a.get('headline'))
        if _is_default(a):
            label, bg, fg = layout.normalize_section(a.get('section'), a.get('headline') or ''), fam[2], fam[3]
        else:
            label, bg, fg = (a.get('section') or brand['name']), brand.get('color') or '#2A2C29', '#FFFFFF'
        sl = prs.slides.add_slide(blank)
        photo = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(1080), px(760))
        photo.fill.solid(); photo.fill.fore_color.rgb = rgb('#D9D7D0'); photo.line.fill.background()
        tf = photo.text_frame; tf.text = 'FOTO'; tf.paragraphs[0].runs[0].font.size = Pt(28); tf.paragraphs[0].runs[0].font.color.rgb = rgb('#9A988F')
        band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, px(760), px(1080), px(590))
        band.fill.solid(); band.fill.fore_color.rgb = rgb('#FFFFFF'); band.line.fill.background()
        pill = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(60), px(700), px(max(180, 22 * len(label) + 60)), px(64))
        pill.fill.solid(); pill.fill.fore_color.rgb = rgb(bg); pill.line.fill.background()
        ptf = pill.text_frame; ptf.text = label.upper()
        r = ptf.paragraphs[0].runs[0]; r.font.bold = True; r.font.size = Pt(20); r.font.color.rgb = rgb(fg); r.font.name = 'Montserrat'
        head = (a.get('headline') or '').strip(); sub = (a.get('subtitle') or '').strip()
        size = 44 if len(head) <= 70 else 38 if len(head) <= 110 else 32 if len(head) <= 160 else 26  # todo entra, sin recortar
        hb = sl.shapes.add_textbox(px(60), px(790), px(960), px(360)); htf = hb.text_frame; htf.word_wrap = True; htf.vertical_anchor = MSO_ANCHOR.TOP
        htf.text = head; hr = htf.paragraphs[0].runs[0]; hr.font.bold = True; hr.font.size = Pt(size); hr.font.name = 'Montserrat'; hr.font.color.rgb = rgb('#1B1D1A')
        if sub:
            sb = sl.shapes.add_textbox(px(60), px(1160), px(960), px(150)); stf = sb.text_frame; stf.word_wrap = True
            stf.text = sub; sr = stf.paragraphs[0].runs[0]; sr.font.size = Pt(20 if len(sub) <= 180 else 16); sr.font.name = 'Montserrat'; sr.font.color.rgb = rgb('#4A4A45')
        brand_box = sl.shapes.add_textbox(px(60), px(1300), px(960), px(40))
        brand_box.text_frame.text = brand['name']; br = brand_box.text_frame.paragraphs[0].runs[0]; br.font.size = Pt(14); br.font.bold = True; br.font.color.rgb = rgb(brand.get('color') or '#3A33E6')
    out = BytesIO(); prs.save(out)
    return out.getvalue()


LOTE_TEMPLATES = {'infolinense': os.getenv('CANVA_LOTE_TEMPLATE_ID', '').strip() or 'EAHZHiNN0bA'}
LOTE_FAMILIES = 12   # páginas por sección en la plantilla de lote: p, p+12, p+24, p+36
LOTE_SLOTS = 4


def lote_plan(articles):
    """Reparte las noticias en huecos de la plantilla de lote según su sección.
    Devuelve una lista de lotes; cada lote es {número de página: noticia}."""
    lotes = []
    for a in articles:
        p = int(_page_for(a) or 1)
        p = p if 1 <= p <= LOTE_FAMILIES else 1
        for lote in lotes:
            free = [p + LOTE_FAMILIES * k for k in range(LOTE_SLOTS) if p + LOTE_FAMILIES * k not in lote]
            if free:
                lote[free[0]] = a; break
        else:
            lotes.append({p: a})
    return lotes


def _lote_fields(lote, schema):
    fields = {}
    for n in range(1, LOTE_FAMILIES * LOTE_SLOTS + 1):
        a = lote.get(n)
        t = design_fields(a, strict=False) if a else {}
        for k in ('HEADLINE', 'SUMMARY', 'SECTION'):
            fields['%s_%d' % (k, n)] = {'type': 'text', 'text': (t.get(k) or '').strip() or BLANK}
    for k in ('HEADLINE', 'SUMMARY', 'SECTION'):   # restos sin número de la plantilla original
        if schema.get(k, {}).get('type') == 'text':
            fields[k] = {'type': 'text', 'text': BLANK}
    return {k: v for k, v in fields.items() if not schema or schema.get(k, {}).get('type') == 'text'}


def _lote_with_template(articles, template_id, progress, start=0):
    from datetime import datetime
    schema = _api('GET', '/brand-templates/' + template_id + '/dataset').get('dataset', {})
    out, done = [], start
    lotes = lote_plan(articles)
    stamp = datetime.now().strftime('%d/%m/%Y %H:%M')
    for i, lote in enumerate(lotes, 1):
        progress(done, 'Rellenando lote %d de %d' % (i, len(lotes)))
        job = _api('POST', '/autofills', json={'type': 'create_from_brand_template', 'brand_template_id': template_id,
                                              'title': 'Lote %d · %s' % (i, stamp), 'data': _lote_fields(lote, schema)})
        design = _design_from_job(_wait('/autofills/' + job['job']['id'], job))
        if not design.get('id'):
            raise ValueError('Canva no devolvió el diseño del lote %d' % i)
        url = design.get('url') or (design.get('urls') or {}).get('edit_url') or 'https://www.canva.com/design/%s/edit' % design['id']
        done += len(lote)
        out.append({'url': url, 'design_id': design['id'], 'pages': len(lote)})
    return out


def _lote_pptx(articles, progress, start=0):
    from datetime import datetime
    data = _batch_pptx(articles)
    progress(start + len(articles), 'Abriendo en Canva')
    title = 'Lote de noticias · ' + datetime.now().strftime('%d/%m/%Y %H:%M')
    meta = json.dumps({'title_base64': base64.b64encode(title.encode()).decode(),
                       'mime_type': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'})
    job = _api('POST', '/imports', data=data, headers={'Content-Type': 'application/octet-stream', 'Import-Metadata': meta})
    done = _wait('/imports/' + job['job']['id'], job)
    designs = (done.get('result') or {}).get('designs') or done.get('designs') or []
    design = designs[0] if designs else {}
    if not design.get('id'):
        raise ValueError('Canva no devolvió el diseño importado')
    url = design.get('url') or (design.get('urls') or {}).get('edit_url') or 'https://www.canva.com/design/%s/edit' % design['id']
    return [{'url': url, 'design_id': design['id'], 'pages': len(articles)}]


def batch_design(articles, progress=None):
    """Lote para Canva con TU plantilla: la plantilla «Lote InfoLinense» tiene 4 huecos por sección
    (48 páginas). Cada noticia va al hueco de su sección; si una sección pasa de 4, se abre otro lote.
    Medios sin plantilla de lote: un diseño importado con todas las páginas."""
    if not ready() or not connected():
        raise ValueError('Conecta tu cuenta de Canva desde Ajustes')
    from . import brands
    progress = progress or (lambda d, ph: None)
    groups = {}
    for a in articles:
        groups.setdefault(brands.valid(a.get('brand')), []).append(a)
    lotes, done = [], 0
    for slug, arts in groups.items():
        tpl = LOTE_TEMPLATES.get(slug)
        lotes += _lote_with_template(arts, tpl, progress, done) if tpl else _lote_pptx(arts, progress, done)
        done += len(arts)
    progress(done, 'Listo')
    first = lotes[0] if lotes else {}
    return {'url': first.get('url', ''), 'design_id': first.get('design_id', ''), 'pages': len(articles),
            'merged': True, 'lotes': lotes, 'designs': [], 'errors': []}
