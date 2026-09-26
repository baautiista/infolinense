import os, json, threading, time, ipaddress
from datetime import datetime, timezone
from urllib.parse import urlparse
from zoneinfo import ZoneInfo
from io import BytesIO
from pathlib import Path
from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Depends
from fastapi.responses import FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from pptx import Presentation
from . import db, sources, pipeline, renderer, publishers, photos, canva
from .auth import login, require_auth
from .config import BASE_DIR, RENDER_DIR, TEMPLATE_DIR, UPLOAD_DIR, SCAN_INTERVAL_MINUTES, AUTO_PIPELINE, OPENAI_API_KEY, ADMIN_PASSWORD, JWT_SECRET, PUBLISH_MODE, CORS_ORIGINS, PUBLIC_BASE_URL

app=FastAPI(title='InfoLinense Desk',version='2.1.0')
app.add_middleware(CORSMiddleware,allow_origins=CORS_ORIGINS,allow_credentials=False,allow_methods=['*'],allow_headers=['Authorization','Content-Type'])
db.init_db()

class LoginIn(BaseModel): password:str
class SourceIn(BaseModel):
    name:str
    url:str
    kind:str='rss'
    priority:int=50
    official:bool=False
    local_scope:bool=False
class EditArticle(BaseModel):
    section:str|None=None; headline:str|None=None; subtitle:str|None=None; body:str|None=None; graphic_summary:str|None=None
class PhotoChoice(BaseModel): url:str; source:str=''; license:str=''; author:str=''

@app.post('/api/auth/login')
def auth_login(body:LoginIn): return {'token':login(body.password)}
@app.get('/api/health')
def health():
    return {'ok':True,'version':'2.1.0','auth_configured':bool(ADMIN_PASSWORD and JWT_SECRET),'openai_configured':bool(OPENAI_API_KEY),'draft_mode':'ai' if OPENAI_API_KEY else 'source_brief','publish_mode':PUBLISH_MODE}
@app.get('/api/capabilities',dependencies=[Depends(require_auth)])
def capabilities():
    return {'real_photo_only':True,'ai_image_generation':False,'max_social_chars':2200,'auto_pipeline':AUTO_PIPELINE,'public_base_url':PUBLIC_BASE_URL,'canva_configured':canva.ready(),'canva_connected':canva.connected(),'canva_redirect_uri':canva.callback_url(),'primary_template':'canva' if canva.ready() else 'pptx','canva_template_url':'https://www.canva.com/brand/brand-templates/'+canva.TEMPLATE_ID if canva.TEMPLATE_ID else None}

@app.get('/api/canva/connect',dependencies=[Depends(require_auth)])
def connect_canva():
    try: return {'url':canva.authorization_url()}
    except ValueError as e: raise HTTPException(400,str(e))

@app.get('/api/canva/template',dependencies=[Depends(require_auth)])
def canva_template_status():
    try: return canva.template_status()
    except ValueError as e: raise HTTPException(400,str(e))

@app.get('/api/canva/permissions',dependencies=[Depends(require_auth)])
def canva_permissions():
    try: return canva.permission_status()
    except ValueError as e: raise HTTPException(400,str(e))

@app.get('/api/canva/callback')
def canva_callback(code:str='',state:str='',error:str=''):
    if error: raise HTTPException(400,'Canva no autorizó la conexión')
    try: canva.complete(code,state)
    except ValueError as e: raise HTTPException(400,str(e))
    return RedirectResponse('https://infolinense-desk.lovable.app/?canva=connected',status_code=303)

@app.post('/api/articles/{aid}/canva',dependencies=[Depends(require_auth)])
def create_canva_design(aid:int):
    a=db.row('SELECT * FROM articles WHERE id=?',(aid,))
    if not a: raise HTTPException(404)
    try: return canva.create_design(a)
    except ValueError as e: raise HTTPException(400,str(e))
@app.get('/api/dashboard',dependencies=[Depends(require_auth)])
def dashboard():
    return {'counts':{'new':db.row("SELECT COUNT(*) n FROM candidates WHERE status='new'")['n'],
                      'researched':db.row("SELECT COUNT(*) n FROM candidates WHERE status='researched'")['n'],
                      'draft':db.row("SELECT COUNT(*) n FROM candidates WHERE status='draft'")['n'],
                      'review':db.row("SELECT COUNT(*) n FROM articles WHERE status='review_ready'")['n'],
                      'approved':db.row("SELECT COUNT(*) n FROM articles WHERE status='approved'")['n'],
                      'published':db.row("SELECT COUNT(*) n FROM articles WHERE status='published'")['n']},
            'latest':db.rows("SELECT * FROM candidates WHERE status IN ('new','researched','draft') ORDER BY id DESC LIMIT 15"),
            'activity':db.rows('SELECT * FROM activity ORDER BY id DESC LIMIT 12')}
@app.get('/api/radar/stats',dependencies=[Depends(require_auth)])
def radar_stats():
    midnight=datetime.now(ZoneInfo('Europe/Madrid')).replace(hour=0,minute=0,second=0,microsecond=0)
    start=midnight.astimezone(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')
    entries=db.rows('SELECT published_at FROM candidates WHERE created_at>=?',(start,))
    dated=[sources.publication_datetime(x['published_at']) for x in entries]
    confirmed=sum(1 for date in dated if date and date>=midnight.astimezone(timezone.utc))
    undated=sum(1 for date in dated if date is None)
    return {'date':midnight.date().isoformat(),'timezone':'Europe/Madrid','target':50,'found':confirmed,
            'discovered_today':len(entries),'undated_today':undated,
            'source_count':db.row('SELECT COUNT(*) n FROM sources WHERE active=1')['n'],
            'sources':db.rows('''SELECT s.id,s.name,s.kind,s.active,s.last_checked_at,s.last_success_at,
                       s.last_error,s.items_seen,s.items_added,
                       (SELECT COUNT(*) FROM candidates c WHERE c.source_id=s.id AND c.created_at>=?) today
                       FROM sources s ORDER BY s.priority DESC''',(start,))}
@app.post('/api/scan',dependencies=[Depends(require_auth)])
def scan():
    r=sources.scan_all()
    if AUTO_PIPELINE and not r.get('busy'): r['pipeline']=pipeline.auto_process(r['added'])
    return r
@app.get('/api/candidates',dependencies=[Depends(require_auth)])
def candidates(status:str='new'):
    if status=='pending':
        return db.rows("SELECT * FROM candidates WHERE status IN ('new','researched','draft') ORDER BY score DESC,id DESC")
    return db.rows('SELECT * FROM candidates WHERE status=? ORDER BY score DESC,id DESC',(status,))
@app.post('/api/candidates/{cid}/investigate',dependencies=[Depends(require_auth)])
def investigate(cid:int):
    try: return {'ok':True,'research':pipeline.investigate_candidate(cid)}
    except ValueError as e: raise HTTPException(400,str(e))
    except RuntimeError as e: raise HTTPException(404,str(e))
@app.post('/api/candidates/{cid}/draft',dependencies=[Depends(require_auth)])
def draft_candidate(cid:int):
    try:
        aid=pipeline.draft_candidate(cid)
        return db.row('SELECT * FROM articles WHERE id=?',(aid,))
    except ValueError as e: raise HTTPException(400,str(e))
    except RuntimeError as e: raise HTTPException(404,str(e))
@app.get('/api/candidates/{cid}/draft',dependencies=[Depends(require_auth)])
def get_candidate_draft(cid:int):
    a=db.row("SELECT * FROM articles WHERE candidate_id=? AND status='draft'",(cid,))
    if not a: raise HTTPException(404,'Borrador no disponible')
    return a
@app.post('/api/candidates/{cid}/submit',dependencies=[Depends(require_auth)])
def submit_candidate(cid:int):
    try: return {'id':pipeline.submit_candidate(cid)}
    except ValueError as e: raise HTTPException(400,str(e))
@app.post('/api/candidates/{cid}/archive',dependencies=[Depends(require_auth)])
def archive_candidate(cid:int):
    c=db.row('SELECT id,status FROM candidates WHERE id=?',(cid,))
    if not c: raise HTTPException(404)
    if c['status'] not in ('new','researched','draft'): raise HTTPException(400,'Esta pieza ya está en revisión')
    db.exec_("UPDATE candidates SET status='archived' WHERE id=?",(cid,))
    db.exec_("UPDATE articles SET status='rejected' WHERE candidate_id=? AND status='draft'",(cid,))
    return {'ok':True}
@app.post('/api/candidates/{cid}/prepare',dependencies=[Depends(require_auth)])
def prepare(cid:int,research:bool=False):
    try:
        aid=pipeline.process_candidate(cid,force_research=research); return db.row('SELECT * FROM articles WHERE id=?',(aid,))
    except Exception as e: raise HTTPException(500,str(e))
@app.get('/api/review',dependencies=[Depends(require_auth)])
def review():
    return db.rows("""SELECT a.*,c.title source_title,c.url source_url,c.score,c.source_name,d.url canva_url,d.exported canva_exported FROM articles a LEFT JOIN candidates c ON c.id=a.candidate_id LEFT JOIN canva_designs d ON d.article_id=a.id WHERE a.status IN ('review_ready','approved') ORDER BY a.id DESC""")
@app.get('/api/articles/{aid}',dependencies=[Depends(require_auth)])
def article(aid:int):
    a=db.row('SELECT a.*,c.url source_url,c.score,c.source_name FROM articles a LEFT JOIN candidates c ON c.id=a.candidate_id WHERE a.id=?',(aid,))
    if not a: raise HTTPException(404)
    return a
@app.get('/api/articles/{aid}/kit',dependencies=[Depends(require_auth)])
def kit(aid:int):
    a=article(aid)
    design=db.row('SELECT url,exported FROM canva_designs WHERE article_id=?',(aid,)) or {}
    return {'id':aid,'section':a.get('section'),'headline':a.get('headline'),'text':(a.get('body') or '')[:2200],'chars':len((a.get('body') or '')[:2200]),'image_url':f'/media/render/{aid}.png','status':a.get('status'),'source_url':a.get('source_url'),'image_source':a.get('image_source'),'image_license':a.get('image_license'),'ai_image_suggestion':a.get('ai_image_suggestion') or '','canva_url':design.get('url'),'canva_exported':bool(design.get('exported'))}
@app.put('/api/articles/{aid}',dependencies=[Depends(require_auth)])
def edit_article(aid:int,body:EditArticle):
    original=db.row('SELECT * FROM articles WHERE id=?',(aid,))
    if not original: raise HTTPException(404)
    fields=[]; vals=[]
    for k,v in body.model_dump(exclude_none=True).items():
        if k=='body':
            v=v[:2200]
            fields.append('social_text=?'); vals.append(v)
        fields.append(f'{k}=?'); vals.append(v)
    if fields:
        db.exec_(f"UPDATE articles SET {','.join(fields)},updated_at=CURRENT_TIMESTAMP WHERE id=?",tuple(vals+[aid]))
        db.exec_('UPDATE canva_designs SET exported=0 WHERE article_id=?',(aid,))
        if original['status']=='approved':
            db.exec_("UPDATE articles SET status='review_ready' WHERE id=?",(aid,))
    try: renderer.render_article(aid)
    except Exception as e: db.log('render_error',str(e))
    return db.row('SELECT * FROM articles WHERE id=?',(aid,))
@app.get('/api/articles/{aid}/photos',dependencies=[Depends(require_auth)])
def article_photos(aid:int,refresh:bool=False):
    a=db.row('SELECT a.*,c.title source_title,c.url source_url,c.excerpt,c.source_id FROM articles a LEFT JOIN candidates c ON c.id=a.candidate_id WHERE a.id=?',(aid,))
    if not a: raise HTTPException(404)
    if not refresh:
        try: cached=json.loads(a.get('image_candidates_json') or '[]')
        except Exception: cached=[]
        if cached: return cached
    cand={'title':a.get('source_title') or a.get('headline'),'url':a.get('source_url'),'excerpt':a.get('excerpt'),'source_id':a.get('source_id')}
    imgs=photos.search_real_photos(cand,a.get('section') or 'CIUDAD')
    db.exec_('UPDATE articles SET image_candidates_json=? WHERE id=?',(json.dumps(imgs,ensure_ascii=False),aid)); return imgs
@app.post('/api/articles/{aid}/photo',dependencies=[Depends(require_auth)])
def choose_photo(aid:int,p:PhotoChoice):
    try: local=photos.download_image(p.url)
    except Exception as e: raise HTTPException(400,'No se pudo descargar la foto: '+str(e))
    db.exec_("UPDATE articles SET image_url=?,image_source=?,image_license=?,image_local=?,ai_image_suggestion='',status=CASE WHEN status='approved' THEN 'review_ready' ELSE status END WHERE id=?",(p.url,p.source,p.license,local,aid))
    db.exec_('UPDATE canva_designs SET exported=0 WHERE article_id=?',(aid,))
    try: renderer.render_article(aid)
    except Exception as e: db.log('render_error',str(e))
    return {'ok':True}

@app.post('/api/articles/{aid}/photo/upload',dependencies=[Depends(require_auth)])
async def upload_article_photo(aid:int,file:UploadFile=File(...),license:str=Form(...),source:str=Form('Fotografía propia'),author:str=Form('')):
    if not db.row('SELECT id FROM articles WHERE id=?',(aid,)): raise HTTPException(404)
    if not license.strip() or not source.strip(): raise HTTPException(400,'Indica la licencia y la procedencia de la fotografía')
    data=await file.read(20_000_001)
    if len(data)>20_000_000: raise HTTPException(400,'La fotografía debe pesar menos de 20 MB')
    try:
        from PIL import Image
        from uuid import uuid4
        im=Image.open(BytesIO(data))
        if im.format not in ('JPEG','PNG','WEBP'): raise ValueError('Formato no permitido')
        if im.width*im.height>40_000_000: raise ValueError('La imagen tiene demasiados píxeles')
        im.load()
        dest=UPLOAD_DIR/(uuid4().hex+'.jpg')
        im.convert('RGB').save(dest,'JPEG',quality=92)
    except Exception as e: raise HTTPException(400,'No se pudo leer la fotografía: '+str(e))
    selected={'url':'upload:'+dest.stem,'source':source.strip()[:180],'license':license.strip()[:180],
              'author':author.strip()[:120],'kind':'user_upload','publish_safe':True}
    db.exec_("UPDATE articles SET image_url=?,image_source=?,image_license=?,image_local=?,image_candidates_json=?,status=CASE WHEN status='approved' THEN 'review_ready' ELSE status END,updated_at=CURRENT_TIMESTAMP WHERE id=?",
             (selected['url'],selected['source'],selected['license'],str(dest),json.dumps([selected]),aid))
    db.exec_('UPDATE canva_designs SET exported=0 WHERE article_id=?',(aid,))
    try: renderer.render_article(aid)
    except Exception as e: db.log('render_error',str(e))
    return {'ok':True,'photo':selected}
@app.post('/api/articles/{aid}/render',dependencies=[Depends(require_auth)])
def rerender(aid:int):
    artwork=db.row('SELECT exported FROM canva_designs WHERE article_id=?',(aid,))
    if artwork and artwork['exported']:
        p=RENDER_DIR/f'article_{aid}.png'
        if p.is_file(): return {'path':str(p),'url':f'/media/render/{aid}.png'}
    try: return {'path':renderer.render_article(aid),'url':f'/media/render/{aid}.png'}
    except Exception as e: raise HTTPException(500,str(e))
@app.post('/api/articles/{aid}/approve',dependencies=[Depends(require_auth)])
def approve(aid:int):
    a=db.row('SELECT * FROM articles WHERE id=?',(aid,))
    if not a: raise HTTPException(404)
    if a['status']!='review_ready': raise HTTPException(400,'Solo se puede aprobar una pieza en revisión')
    if not a.get('image_local') or not Path(a['image_local']).is_file() or not a.get('image_license'):
        raise HTTPException(400,'Selecciona una fotografía real con licencia antes de aprobar')
    try: candidates=json.loads(a.get('image_candidates_json') or '[]')
    except (TypeError,ValueError): candidates=[]
    selected=next((p for p in candidates if p.get('url')==a.get('image_url')),None)
    if not selected or not selected.get('publish_safe'):
        raise HTTPException(400,'No se ha verificado el permiso de reutilización de esta fotografía')
    if canva.ready() and canva.connected():
        artwork=db.row('SELECT exported FROM canva_designs WHERE article_id=?',(aid,))
        if not artwork or not artwork['exported']:
            raise HTTPException(400,'Crea y exporta primero el diseño final con la plantilla de Canva')
    db.exec_("UPDATE articles SET status='approved',updated_at=CURRENT_TIMESTAMP WHERE id=?",(aid,))
    return {'ok':True}
@app.post('/api/articles/{aid}/reject',dependencies=[Depends(require_auth)])
def reject(aid:int): db.exec_("UPDATE articles SET status='rejected',updated_at=CURRENT_TIMESTAMP WHERE id=?",(aid,)); return {'ok':True}
@app.post('/api/articles/{aid}/research-more',dependencies=[Depends(require_auth)])
def research_more(aid:int):
    a=db.row('SELECT candidate_id FROM articles WHERE id=?',(aid,))
    if not a: raise HTTPException(404)
    db.exec_('DELETE FROM articles WHERE id=?',(aid,)); db.exec_("UPDATE candidates SET status='new' WHERE id=?",(a['candidate_id'],))
    new_aid=pipeline.process_candidate(a['candidate_id'],force_research=True)
    return db.row('SELECT * FROM articles WHERE id=?',(new_aid,))
@app.post('/api/articles/{aid}/publish',dependencies=[Depends(require_auth)])
def publish(aid:int):
    a=db.row('SELECT * FROM articles WHERE id=?',(aid,))
    if not a: raise HTTPException(404)
    if a['status']!='approved': raise HTTPException(400,'Primero debes aprobar la noticia')
    if canva.ready() and canva.connected():
        artwork=db.row('SELECT exported FROM canva_designs WHERE article_id=?',(aid,))
        if not artwork or not artwork['exported']:
            raise HTTPException(400,'Regenera la imagen final en Canva antes de publicar')
    try:
        url=publishers.publish(a); db.exec_("UPDATE articles SET status='published',publish_url=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(url,aid)); return {'ok':True,'url':url}
    except Exception as e: raise HTTPException(400,str(e))
@app.get('/api/sources',dependencies=[Depends(require_auth)])
def list_sources(): return db.rows('SELECT * FROM sources ORDER BY priority DESC')
def validate_source(s):
    url=urlparse(s.url.strip())
    host=(url.hostname or '').lower()
    if url.scheme!='https' or not host or url.username or url.password or host in {'localhost','localhost.localdomain'} or host.endswith(('.local','.internal')):
        raise HTTPException(400,'La fuente debe tener una URL HTTPS pública')
    try:
        if not ipaddress.ip_address(host).is_global: raise HTTPException(400,'No se permiten direcciones privadas')
    except ValueError: pass
    if s.kind not in ('rss','html','bop','procurement','social') or not 0<=s.priority<=100 or not s.name.strip() or len(s.name)>120:
        raise HTTPException(400,'Revisa el nombre, tipo y prioridad de la fuente')
    if s.kind=='social' and host not in ('facebook.com','www.facebook.com','instagram.com','www.instagram.com'):
        raise HTTPException(400,'Las fuentes sociales aceptan páginas de Facebook o Instagram')
    if s.kind=='social' and not url.path.strip('/'):
        raise HTTPException(400,'Indica la página o cuenta pública')
    if s.kind=='procurement' and host!='contratos.gobierto.es':
        raise HTTPException(400,'El extractor de licitaciones corresponde a Gobierto')
    if s.kind=='bop' and host not in ('bopcadiz.es','www.bopcadiz.es'):
        raise HTTPException(400,'El extractor del BOP corresponde a Cádiz')
    return s.name.strip(),s.url.strip(),s.kind,s.priority,int(s.official),int(s.local_scope)
@app.post('/api/sources',dependencies=[Depends(require_auth)])
def add_source(s:SourceIn):
    values=validate_source(s)
    try: return {'id':db.exec_('INSERT INTO sources(name,url,kind,priority,official,local_scope) VALUES(?,?,?,?,?,?)',values)}
    except Exception as e: raise HTTPException(400,str(e))
@app.put('/api/sources/{sid}',dependencies=[Depends(require_auth)])
def edit_source(sid:int,s:SourceIn):
    if not db.row('SELECT id FROM sources WHERE id=?',(sid,)): raise HTTPException(404)
    values=validate_source(s)
    try:
        db.exec_('''UPDATE sources SET name=?,url=?,kind=?,priority=?,official=?,local_scope=?,
                   last_checked_at=NULL,last_error=NULL WHERE id=?''',(*values,sid))
    except Exception as e: raise HTTPException(400,str(e))
    return {'ok':True}
@app.post('/api/sources/{sid}/toggle',dependencies=[Depends(require_auth)])
def toggle_source(sid:int):
    if not db.row('SELECT id FROM sources WHERE id=?',(sid,)): raise HTTPException(404)
    db.exec_('UPDATE sources SET active=CASE active WHEN 1 THEN 0 ELSE 1 END WHERE id=?',(sid,))
    return {'ok':True}
@app.post('/api/sources/{sid}/check',dependencies=[Depends(require_auth)])
def check_source(sid:int):
    if not db.row('SELECT id FROM sources WHERE id=?',(sid,)): raise HTTPException(404)
    return sources.probe_source(sid)
@app.get('/api/templates',dependencies=[Depends(require_auth)])
def list_templates(): return db.rows('SELECT * FROM templates ORDER BY id')
@app.post('/api/templates/{tid}/activate',dependencies=[Depends(require_auth)])
def activate_template(tid:int):
    if not db.row('SELECT id FROM templates WHERE id=?',(tid,)): raise HTTPException(404)
    c=db.conn()
    try:
        c.execute('UPDATE templates SET active=0')
        c.execute('UPDATE templates SET active=1 WHERE id=?',(tid,))
        c.commit()
    finally: c.close()
    return {'ok':True}
@app.post('/api/templates/upload',dependencies=[Depends(require_auth)])
async def upload_template(file:UploadFile=File(...),name:str=Form('Plantilla personalizada')):
    if not file.filename.lower().endswith('.pptx'): raise HTTPException(400,'Solo PPTX')
    data=await file.read(20_000_001)
    if len(data)>20_000_000: raise HTTPException(400,'El PPTX debe pesar menos de 20 MB')
    try:
        prs=Presentation(BytesIO(data))
        if len(prs.slides) not in (1,12): raise ValueError('La plantilla debe tener una diapositiva o 12 variantes')
        if not all('{{HEADLINE}}' in ' '.join(sh.text for sh in slide.shapes if sh.has_text_frame) for slide in prs.slides):
            raise ValueError('Cada diapositiva debe incluir el texto {{HEADLINE}}')
    except Exception as e: raise HTTPException(400,f'Plantilla no válida: {e}')
    from uuid import uuid4
    dest=TEMPLATE_DIR/(uuid4().hex+'.pptx'); dest.write_bytes(data)
    sm={k:(i if len(prs.slides)==12 else 1) for i,k in enumerate(('URBANISMO','CIUDAD','GIBRALTAR','SUCESOS','CULTURA','DEPORTES','COMERCIO','MEDIO AMBIENTE','POLÍTICA','SOCIEDAD','PATRIMONIO','AGENDA'),1)}
    c=db.conn()
    try:
        c.execute('UPDATE templates SET active=0')
        cur=c.execute('INSERT INTO templates(name,path,format,slide_map_json,active) VALUES(?,?,?,?,1)',(name,str(dest),'4:5',json.dumps(sm,ensure_ascii=False)))
        c.commit()
        return {'id':cur.lastrowid,'active':True}
    finally: c.close()

@app.get('/media/render/{aid}.png',dependencies=[Depends(require_auth)])
def render_file(aid:int):
    p=RENDER_DIR/f'article_{aid}.png'
    if not p.exists(): raise HTTPException(404)
    return FileResponse(p,media_type='image/png',headers={'Cache-Control':'no-store'})
app.mount('/static',StaticFiles(directory=BASE_DIR/'static'),name='static')
@app.get('/')
def index(): return FileResponse(BASE_DIR/'static'/'index.html')

def scheduler_loop():
    time.sleep(8)
    while True:
        try:
            r=sources.scan_all()
            if AUTO_PIPELINE and not r.get('busy'): pipeline.auto_process(r['added'])
        except Exception as e: db.log('scheduler_error',str(e))
        time.sleep(max(15,SCAN_INTERVAL_MINUTES)*60)
@app.on_event('startup')
def startup():
    if os.getenv('DISABLE_SCHEDULER','false').lower() not in {'1','true','yes'}: threading.Thread(target=scheduler_loop,daemon=True).start()
